"""Local, values-only XLSX ingestion; the analysis runner continues to use CSV."""
import csv
from datetime import date, time
import io
import re
import zipfile

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLUMNS = 200


def xlsx_csv(raw, sheet=None):
    """Return CSV bytes plus sheet metadata, or a sheet-choice response. Never execute formulas."""
    if len(raw) > MAX_BYTES:
        raise ValueError('Dateien dürfen höchstens 20 MB groß sein.')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist()) > MAX_EXPANDED:
                raise ValueError('Die entpackte XLSX-Datei ist zu groß (maximal 100 MB).')
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except ImportError:
        raise ValueError('Für XLSX fehlt openpyxl. Bitte Einrichtung.cmd erneut ausführen.') from None
    except ValueError:
        raise
    except Exception:
        raise ValueError('Die Datei ist keine lesbare XLSX-Arbeitsmappe. Bitte als .xlsx ohne Kennwort speichern.') from None
    try:
        sheets = workbook.sheetnames
        if sheet is None and len(sheets) > 1:
            return None, {'requires_sheet': True, 'sheets': sheets}
        selected = sheet if sheet is not None else sheets[0]
        if selected not in sheets:
            raise ValueError('Das gewählte Tabellenblatt existiert nicht.')
        ws = workbook[selected]
        if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLUMNS:
            raise ValueError('XLSX-Blätter dürfen höchstens 100000 Zeilen und 200 Spalten umfassen.')
        output = io.StringIO(newline='')
        writer = csv.writer(output, delimiter=';', lineterminator='\n')
        width = None
        for index, cells in enumerate(ws.iter_rows(), 1):
            if index > MAX_ROWS or len(cells) > MAX_COLUMNS:
                raise ValueError('Das Tabellenblatt überschreitet die Importgrenzen.')
            values = []
            for cell in cells:
                if cell.data_type in ('f', 'e'):
                    raise ValueError(f'Zelle {cell.coordinate}: Formel oder Excel-Fehler. Bitte in einer Kopie durch Werte ersetzen.')
                value = cell.value
                if value is None:
                    text = ''
                elif isinstance(value, (date, time)):
                    text = value.isoformat()
                elif isinstance(value, int) and not isinstance(value, bool) and re.fullmatch(r'0+', cell.number_format):
                    text = format(value, '0' + str(len(cell.number_format)) + 'd')
                else:
                    text = str(value)
                values.append(text)
            if width is None:
                while values and values[-1] == '':
                    values.pop()
                if len(values) < 2 or any(not h.strip() for h in values) or len(set(values)) != len(values):
                    raise ValueError('Die erste Excel-Zeile muss eindeutige, ausgefüllte Spaltennamen enthalten. Titelzeilen vorher entfernen.')
                width = len(values)
            else:
                if any(values[width:]):
                    raise ValueError(f'Excel-Zeile {index}: Daten in einer Spalte ohne Überschrift.')
                if not any(values):
                    continue
                values = values[:width] + [''] * max(0, width-len(values))
            writer.writerow(values)
            if output.tell() > MAX_BYTES:
                raise ValueError('Die eingelesenen Tabellenwerte sind zu groß (maximal 20 MB als CSV).')
        result = output.getvalue().encode('utf-8')
        if len(result) > MAX_BYTES:
            raise ValueError('Die eingelesenen Tabellenwerte sind zu groß (maximal 20 MB als CSV).')
        return result, {'sheet': selected, 'sheets': sheets, 'format': 'xlsx'}
    finally:
        workbook.close()
