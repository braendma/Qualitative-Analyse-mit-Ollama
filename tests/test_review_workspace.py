import base64
import copy
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import yaml

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from local_app import App
from review_workspace import decisions_xlsx, compare_books
from codebook_refinement import propose, validate_proposals
from runtime_support import atomic_json, file_hash
from progress_events import update_progress, request_event
from setup_checks import check_setup, test_local_model
from telegram_notifications import Telegram
from test_extensions import SEGS, BOOK, VERIFY, fake_blind, PARAMS
from multi_label_core import blind_code_units, calculate_set_agreement
from review_queue import build_queue
from coding_validation_common import load_segments


def fixture(directory):
    app=App(directory);pid=app.create('Synthetic review')['id']
    seg='segment_id;PassageID;Dokumentname;Code;Segment\nS1;U1;P1;A;Gleicher Text\nS2;U1;P1;B;Gleicher Text\nS3;U2;P2;A;Anderer Text\n'
    book='Kategorie;Unterkategorie;Ausprägung;Facette;Definition;Ankerbeispiel\n'+''.join(f'{c};;;;{c};{c}\n' for c in ['A','B','C'])
    for kind,text in [('segments',seg),('codebook',book)]:app.upload(pid,kind,kind+'.csv',base64.b64encode(text.encode()).decode())
    settings={'model':'mock','columns':{'segment_id':'segment_id','unit_id':'PassageID','person':'Dokumentname','code':'Code','segment':'Segment'},
              'book_columns':dict(zip(('kategorie','unterkategorie','auspraegung','facette','definition','ankerbeispiel'),('Kategorie','Unterkategorie','Ausprägung','Facette','Definition','Ankerbeispiel'))),
              'modules':['review_queue'],'label_mode':'multi_label'}
    checked=app.save(pid,settings);cfg=app.project_dir(pid)/'revisions'/app.project(pid)['revision']/'config.yaml'
    _,blind=blind_code_units(SEGS,BOOK,{}, {},PARAMS,llm=fake_blind)
    _,agreement=calculate_set_agreement(SEGS,BOOK,VERIFY,blind)
    queue=build_queue(SEGS,BOOK,agreement,VERIFY,blind,{'befunde':[]})
    jid='a'*20;folder=app.project_dir(pid)/'jobs'/jid;run=folder/'runs'/'synthetic'
    atomic_json(run/'review_queue.json',queue)
    atomic_json(run/'workflow_manifest.json',{'status':'success','completed_steps':['review_queue'],'current_module':None})
    atomic_json(folder/'job.json',{'id':jid,'created':1,'config':str(cfg),'status':'success','modules':checked['modules']})
    return app,pid,jid,queue,cfg


def decision(**changes):
    return {'case_id':'unit:U1','decision':'custom','final_codes':['A','C'],'note':'Die Aussage beschreibt zwei Aspekte.','reviewer':'Testperson',**changes}


class ReviewTests(unittest.TestCase):
    def test_review_http_requires_auth_and_preserves_conflicting_draft(self):
        import http.client
        import threading
        from local_app import make_server
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);server=make_server(app)
            thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
            def request(method,path,data=None,authorized=True):
                connection=http.client.HTTPConnection('127.0.0.1',server.server_port,timeout=5)
                headers={'Content-Type':'application/json'}
                if authorized:headers['X-App-Token']=server.token
                connection.request(method,path,json.dumps(data) if data is not None else None,headers)
                response=connection.getresponse();result=(response.status,response.read());connection.close();return result
            query=f'?project={pid}&job={jid}'
            try:
                self.assertEqual(request('GET','/api/review'+query,authorized=False)[0],403)
                data={'project':pid,'job':jid,'decision':decision(),'revision':0}
                self.assertEqual(request('POST','/api/review-save',data)[0],200)
                self.assertEqual(request('POST','/api/review-save',data)[0],400)
                status,body=request('GET','/api/review-export'+query+'&format=xlsx')
                self.assertEqual(status,200);self.assertTrue(body.startswith(b'PK'))
                status,body=request('GET','/api/review-export'+query+'&format=json')
                self.assertEqual(json.loads(body)['revision'],1)
                self.assertEqual(request('POST','/api/followup-prepare',{'project':pid,'job':jid,'revision':True})[0],400)
                self.assertEqual(len(app.jobs(pid)),1)
            finally:server.shutdown();server.server_close();thread.join()

    def test_category_versions_include_old_job_source_but_not_failed_draft(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp)
            (cfg.parent/'settings.json').unlink()  # Simulate a project from the previous app version.
            invalid=copy.deepcopy(app.project(pid)['settings']);invalid['columns']['person']='segment_id'
            with self.assertRaises(ValueError):app.save(pid,invalid)
            self.assertEqual([v['id'] for v in app.category_versions(pid)],[cfg.parent.name])

    def test_autosave_draft_restart_history_and_concurrent_conflict(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp)
            source=file_hash(app.artifact(pid,jid,'review_queue.json'))
            first=app.save_review(pid,jid,decision(note='',reviewer=''),0)
            self.assertEqual(first['summary']['critical_open'],1)
            reopened=App(tmp);self.assertEqual(reopened.review(pid,jid)['draft']['revision'],1)
            with self.assertRaisesRegex(ValueError,'anderen Fenster'):reopened.save_review(pid,jid,decision(),0)
            second=reopened.save_review(pid,jid,decision(),1)
            self.assertEqual(second['summary']['critical_open'],0)
            self.assertEqual(len(list((app.project_dir(pid)/'jobs'/jid/'review/versions').glob('*.json'))),2)
            self.assertEqual(file_hash(app.artifact(pid,jid,'review_queue.json')),source)
            with self.assertRaises(ValueError):app.save_review(pid,jid,decision(final_codes=['unbekannt']),2)
            self.assertEqual(app.review(pid,jid)['draft']['revision'],2)

    def test_followup_new_input_keeps_original_and_source_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);before=cfg.read_bytes();old=yaml.safe_load(before);input_before=file_hash(old['paths']['input_csv'])
            with self.assertRaisesRegex(ValueError,'kritischen'):app.followup_preview(pid,jid)
            app.save_review(pid,jid,decision(),0)
            with self.assertRaisesRegex(ValueError,'geändert'):app.prepare_followup(pid,jid,0)
            preview=app.followup_preview(pid,jid);self.assertEqual(preview['changed'],1)
            with patch.object(app,'start',side_effect=AssertionError('Preparation must not launch')):
                prepared=app.prepare_followup(pid,jid,1)
            newpath=app.project_dir(pid)/'revisions'/prepared['project']['revision']/'config.yaml'
            new=yaml.safe_load(newpath.read_text(encoding='utf-8'));segments=load_segments(new['paths']['input_csv'],new['columns'])
            self.assertEqual([s.human_code for s in segments],['A','C','A'])
            self.assertEqual({s.text for s in segments},{s.text for s in SEGS})
            self.assertFalse(new['coding_agreement']['independent_units_confirmed'])
            self.assertEqual(new['review_provenance']['parent_job'],jid)
            self.assertEqual(cfg.read_bytes(),before);self.assertEqual(file_hash(old['paths']['input_csv']),input_before)
            self.assertEqual(len(app.jobs(pid)),1)
            # Save from the UI keeps the explicit follow-up lineage.
            app.save(pid,prepared['project']['settings'])
            saved=app.project_dir(pid)/'revisions'/app.project(pid)['revision']/'config.yaml'
            self.assertEqual(yaml.safe_load(saved.read_text())['review_provenance']['parent_job'],jid)

    def test_uncoded_cases_require_explicit_exclusion(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);app.save_review(pid,jid,decision(final_codes=[]),0)
            self.assertEqual(app.followup_preview(pid,jid)['uncoded'],1)
            with self.assertRaisesRegex(ValueError,'ausdrücklich'):app.prepare_followup(pid,jid,1)
            p=app.prepare_followup(pid,jid,1,True)['project']
            self.assertEqual(p['review_provenance']['excluded_cases'],['unit:U1'])
            self.assertEqual(p['uploads']['segments']['count'],1)

    def test_xlsx_preserves_literals_notes_and_full_text_or_explicit_error(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);app.save_review(pid,jid,decision(note='=HYPERLINK("https://example.invalid")'),0)
            payload=app.review(pid,jid)['draft'];queue['cases'][0]['text']='=1+1\nOriginaltext'
            wb=load_workbook(io.BytesIO(decisions_xlsx(queue,payload)))
            ws=wb.active;self.assertEqual(ws['C2'].value,'=1+1\nOriginaltext');self.assertEqual(ws['C2'].data_type,'s')
            self.assertEqual(ws['H2'].data_type,'s');self.assertTrue(ws['J2'].value=='ja');self.assertEqual(ws.max_row,3)
            queue['cases'][0]['text']='x'*32768
            with self.assertRaisesRegex(ValueError,'JSON'):decisions_xlsx(queue,payload)

    def test_categories_compare_even_if_removed_code_blocks_next_run(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);project=app.project(pid)
            book='Kategorie;Unterkategorie;Ausprägung;Facette;Definition;Ankerbeispiel\nA;;;;New definition;A\nC;;;;C;C\nD;;;;New category;D\n'
            app.upload(pid,'codebook','new.csv',base64.b64encode(book.encode()).decode())
            diff=app.compare_categories(pid,project['settings'])
            self.assertEqual([c['code'] for c in diff['added']],['D']);self.assertEqual([c['code'] for c in diff['removed']],['B'])
            self.assertEqual([c['code'] for c in diff['changed']],['A']);self.assertEqual(diff['affected_count'],3)
            with self.assertRaisesRegex(ValueError,'fehlen'):app.save(pid,project['settings'])
            self.assertIsNone(app.project(pid)['revision'])

    def test_refinement_source_snapshot_and_no_automatic_category_changes(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);app.save_review(pid,jid,decision(),0)
            project_before=app.project(pid)
            with patch.object(app,'start',return_value={'id':'new'}) as start:
                app.start_refinement(pid,jid,1)
            config=start.call_args.kwargs['prepared_config'];parsed=yaml.safe_load(config.read_text(encoding='utf-8'))
            self.assertEqual([m['id'] for m in parsed['pipeline']['modules']],['codebook_refinement'])
            self.assertEqual(parsed['llm']['host'],'http://localhost:11434');self.assertTrue(Path(parsed['refinement']['queue']).is_file())
            self.assertEqual(app.project(pid),project_before)

    def test_refinement_batches_reject_fabricated_references_and_keep_validated_feedback(self):
        with tempfile.TemporaryDirectory() as tmp:
            app,pid,jid,queue,cfg=fixture(tmp);app.save_review(pid,jid,decision(),0);payload=app.review(pid,jid)['draft'];calls=[]
            def mock(messages,params):
                data=json.loads(messages[1]['content']);calls.append(data)
                self.assertNotIn('reviewer',data['reviewed_cases'][0])
                return json.dumps({'proposals':[{'action':'clarify','affected_codes':['A'],
                    'proposed_categories':[{'levels':['A'],'definition':'Präzise neue Definition'}],
                    'reason':'Abgrenzung anhand der Prüfung','case_ids':['unit:U1'],'limitations':'Nur ein Fall geprüft.'}]})
            result=propose(queue,payload,{'model':'mock','num_ctx':16000,'max_tokens':1000},llm=mock)
            self.assertEqual(result['reviewed_cases'],1);self.assertEqual(result['proposals'][0]['proposal_id'],'P0001')
            self.assertEqual(calls[0]['reviewed_cases'][0]['final_codes'],['A','C'])
            bad=copy.deepcopy(result);bad['proposals'][0]['case_ids']=['invented']
            with self.assertRaises(ValueError):validate_proposals(bad,{'A','B','C'},{'unit:U1'})
            with self.assertRaisesRegex(ValueError,'Kontextbudget'):propose(queue,payload,{'model':'mock','num_ctx':100,'max_tokens':90},llm=mock)

    def test_progress_ignores_text_fields_and_tracks_responses(self):
        with tempfile.TemporaryDirectory() as tmp,patch.dict(os.environ,{'WORKFLOW_PROGRESS_FILE':str(Path(tmp)/'progress.json'),'WORKFLOW_MODULE':'blind_coding'}):
            update_progress(completed=2,total=4,unit='passages',private_text='not allowed')
            request_event(True);request_event()
            content=json.loads((Path(tmp)/'progress.json').read_text())
            self.assertNotIn('private_text',content);self.assertEqual(content['requests'],1);self.assertFalse(content['request_active'])
            self.assertEqual(content['completed'],2)

    def test_setup_does_not_start_model_and_telegram_only_counters(self):
        with tempfile.TemporaryDirectory() as tmp:
            app=App(tmp)
            with patch.object(app,'models',return_value={'models':['mock']}),patch('coding_validation_common.default_llm') as llm:
                self.assertEqual(check_setup(app,'mock')['model_calls'],0);llm.assert_not_called()
                llm.return_value='OK';self.assertTrue(test_local_model(app,'mock')['ok']);self.assertEqual(llm.call_count,1)
            bot=Telegram(tmp);bot.save({'token':'123456:'+('A'*30),'chat_id':'123','enabled':True,'events':['progress'],'persist':False})
            with patch('telegram_notifications.urllib.request.build_opener') as factory:
                opener=factory.return_value;opener.open.return_value.__enter__.return_value.read.return_value=b'{"ok":true}'
                bot.send('progress',1,3,detail={'completed':23,'total':80,'unit':'passages','module':'private name','text':'private text'})
                text=json.loads(opener.open.call_args.args[0].data)['text']
                self.assertIn('23 von 80',text);self.assertNotIn('private',text)


if __name__=='__main__':unittest.main()
